from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\17373\Desktop\cecilia_website\portfolio")
OUT_DIR = ROOT / "cv" / "google_docs"

SOURCES = {
    "Yucheng_Zhang_Master_CV_Google_Docs.docx": ROOT / "cv" / "yczresume.tex",
    "Yucheng_Zhang_SWE_Resume_Google_Docs.docx": ROOT / "cv" / "resume_swe.tex",
    "Yucheng_Zhang_ML_AI_Resume_Google_Docs.docx": ROOT / "cv" / "resume_ml_ai.tex",
    "Yucheng_Zhang_Quant_Finance_Resume_Google_Docs.docx": ROOT / "cv" / "resume_quant_finance.tex",
    "Yucheng_Zhang_Data_Engineer_Resume_Google_Docs.docx": ROOT / "cv" / "resume_data_engineer.tex",
    "Yucheng_Zhang_CTC_SWE_Intern_2027_Google_Docs.docx": ROOT
    / "cv"
    / "ctc_swe_intern_2027"
    / "resume_ctc_swe_intern_2027.tex",
    "Yucheng_Zhang_CTC_Quant_Trading_Intern_2027_Google_Docs.docx": ROOT
    / "cv"
    / "ctc_quant_trading_2027"
    / "resume_ctc_quant_trading_2027.tex",
}


@dataclass
class Entry:
    title: str
    date: str
    org: str
    location: str
    bullets: list[str]


def strip_outer_braces(s: str) -> str:
    s = s.strip()
    if s.startswith("{") and s.endswith("}"):
        return s[1:-1]
    return s


def read_brace_arg(text: str, start: int) -> tuple[str, int]:
    assert text[start] == "{"
    depth = 0
    out = []
    i = start
    while i < len(text):
        ch = text[i]
        if ch == "\\" and i + 1 < len(text):
            out.append(ch)
            i += 1
            out.append(text[i])
        elif ch == "{":
            depth += 1
            if depth > 1:
                out.append(ch)
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return "".join(out), i + 1
            out.append(ch)
        else:
            out.append(ch)
        i += 1
    raise ValueError("Unclosed brace")


def command_args(text: str, command: str, nargs: int) -> list[tuple[list[str], int, int]]:
    results = []
    i = 0
    needle = "\\" + command
    while True:
        idx = text.find(needle, i)
        if idx < 0:
            break
        j = idx + len(needle)
        args = []
        for _ in range(nargs):
            while j < len(text) and text[j].isspace():
                j += 1
            if j >= len(text) or text[j] != "{":
                break
            arg, j = read_brace_arg(text, j)
            args.append(arg)
        if len(args) == nargs:
            results.append((args, idx, j))
        i = j
    return results


def clean_tex(text: str) -> str:
    text = text.replace("\n", " ")
    text = re.sub(r"%.*", "", text)
    text = text.replace(r"\&", "&").replace(r"\%", "%").replace(r"\$", "$")
    text = text.replace(r"\_", "_").replace(r"\#", "#")
    text = text.replace(r"\textbar", "|")
    text = text.replace(r"\times", "x")
    text = text.replace(r"$>$", ">").replace(r"$<$", "<")
    text = text.replace(r"$\times$", "x")
    text = text.replace("``", '"').replace("''", '"')
    text = text.replace("--", "-")
    text = re.sub(r"\\href\{[^{}]*\}\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}", r"\1", text)
    for cmd in ["textbf", "textit", "underline", "small", "footnotesize", "textsc"]:
        text = re.sub(rf"\\{cmd}\{{([^{{}}]*(?:\{{[^{{}}]*\}}[^{{}}]*)*)\}}", r"\1", text)
    text = re.sub(r"\\[A-Za-z]+", "", text)
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def split_sections(tex: str) -> list[tuple[str, str]]:
    matches = list(re.finditer(r"\\section\{([^{}]+)\}", tex))
    sections = []
    for i, m in enumerate(matches):
        name = clean_tex(m.group(1))
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else tex.find(r"\end{document}")
        if end < 0:
            end = len(tex)
        sections.append((name, tex[start:end]))
    return sections


def parse_entries(content: str) -> list[Entry]:
    heads = command_args(content, "resumeSubheading", 4)
    entries: list[Entry] = []
    for n, (args, _start, end) in enumerate(heads):
        next_start = heads[n + 1][1] if n + 1 < len(heads) else len(content)
        chunk = content[end:next_start]
        bullets = [clean_tex(a[0][0]) for a in command_args(chunk, "resumeItem", 1)]
        entries.append(
            Entry(
                title=clean_tex(args[0]),
                date=clean_tex(args[1]),
                org=clean_tex(args[2]),
                location=clean_tex(args[3]),
                bullets=[b for b in bullets if b],
            )
        )
    return entries


def parse_compact_education(content: str) -> tuple[list[tuple[str, str, str]], str | None]:
    rows = []
    row_re = re.compile(r"\\textbf\{([^{}]+)\},\s*(.*?)\s*&\s*(.*?)\\\\", re.S)
    for m in row_re.finditer(content):
        rows.append((clean_tex(m.group(1)), clean_tex(m.group(2)), clean_tex(m.group(3))))
    if not rows:
        for entry in parse_entries(content):
            rows.append((entry.title, entry.org, f"{entry.date}, {entry.location}"))
    coursework = None
    cw = re.search(r"\\textbf\{Relevant (?:Coursework|Courses):\}\s*(.*?)(?:\\vspace|$)", content, re.S)
    if cw:
        coursework = clean_tex(cw.group(1))
    return rows, coursework


def parse_skills(content: str) -> list[tuple[str, str]]:
    pairs = []
    for label, rest in re.findall(r"\\textbf\{([^{}]+)\}\s*([^\\]+)", content):
        value = clean_tex(rest)
        label = clean_tex(label)
        if label and value:
            pairs.append((label.rstrip(":"), value))
    return pairs


def parse_publications(content: str) -> list[str]:
    pubs = []
    for line in content.splitlines():
        if r"\item" not in line:
            continue
        item = clean_tex(re.sub(r"\\item\s*(?:\[[^\]]+\])?", "", line))
        if item:
            pubs.append(item)
    if not pubs:
        pubs = [clean_tex(a[0][0]) for a in command_args(content, "resumeItem", 1)]
    return pubs


def set_font(run, size: float, bold=False, italic=False, color="000000") -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc: Document, text: str = "", size: float = 8.6, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_font(r, size, bold, italic)
    return p


def add_right_tab(p, width_in: float) -> None:
    p.paragraph_format.tab_stops.add_tab_stop(Inches(width_in), WD_TAB_ALIGNMENT.RIGHT)


def add_line_with_right(doc, left: str, right: str, *, size=8.8, bold_left=False, italic_left=False):
    p = paragraph(doc, size=size)
    add_right_tab(p, 7.35)
    r = p.add_run(left)
    set_font(r, size, bold_left, italic_left)
    r = p.add_run("\t" + right)
    set_font(r, size)
    return p


def add_section_heading(doc: Document, name: str, master: bool) -> None:
    p = paragraph(doc)
    p.paragraph_format.space_before = Pt(4 if master else 3)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(name.upper())
    set_font(r, 9.7 if master else 9.5, bold=True)


def add_bullet(doc: Document, text: str, master: bool) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.10)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0.2 if master else 0)
    p.paragraph_format.line_spacing = 1.0
    if p.runs:
        p.runs[0].text = ""
    r = p.add_run(text)
    set_font(r, 8.15 if master else 8.05)


def add_entry(doc: Document, entry: Entry, master: bool) -> None:
    add_line_with_right(doc, entry.title, entry.date, size=8.75 if master else 8.55, bold_left=True)
    add_line_with_right(doc, entry.org, entry.location, size=8.15 if master else 8.0, italic_left=True)
    for bullet in entry.bullets:
        add_bullet(doc, bullet, master)


def add_skills(doc: Document, pairs: list[tuple[str, str]], master: bool) -> None:
    for label, value in pairs:
        p = paragraph(doc, size=8.2 if master else 8.05)
        r = p.add_run(label + ": ")
        set_font(r, 8.2 if master else 8.05, bold=True)
        r = p.add_run(value)
        set_font(r, 8.2 if master else 8.05)


def configure_doc(doc: Document, master: bool) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.47)
    section.right_margin = Inches(0.47)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(8.4 if master else 8.25)


def build_docx(tex_path: Path, out_path: Path, master: bool) -> None:
    tex = tex_path.read_text(encoding="utf-8", errors="replace")
    doc = Document()
    configure_doc(doc, master)

    p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Yucheng (Cecilia) Zhang")
    set_font(r, 14.5 if master else 14.0, bold=True)

    contact = (
        "737-343-4836 | ceciliayczhang@gmail.com | Portfolio | "
        "linkedin.com/in/yczcecilia | github.com/yczhangcc"
    )
    paragraph(doc, contact, size=8.2 if master else 8.0, align=WD_ALIGN_PARAGRAPH.CENTER)

    for name, content in split_sections(tex):
        add_section_heading(doc, name, master)
        if name == "Education":
            rows, coursework = parse_compact_education(content)
            for school, degree, when_where in rows:
                p = paragraph(doc, size=8.2 if master else 8.05)
                add_right_tab(p, 7.35)
                r = p.add_run(school + ", ")
                set_font(r, 8.2 if master else 8.05, bold=True)
                r = p.add_run(degree)
                set_font(r, 8.2 if master else 8.05)
                r = p.add_run("\t" + when_where)
                set_font(r, 8.2 if master else 8.05)
            if coursework:
                p = paragraph(doc, size=7.95 if master else 7.8)
                r = p.add_run("Relevant Coursework: ")
                set_font(r, 7.95 if master else 7.8, bold=True)
                r = p.add_run(coursework)
                set_font(r, 7.95 if master else 7.8)
        elif name == "Technical Skills":
            add_skills(doc, parse_skills(content), master)
        elif name == "Publications":
            for i, pub in enumerate(parse_publications(content), start=1):
                p = paragraph(doc, size=7.9 if master else 7.8)
                p.paragraph_format.left_indent = Inches(0.15)
                p.paragraph_format.first_line_indent = Inches(-0.15)
                r = p.add_run(f"{i}. {pub}")
                set_font(r, 7.9 if master else 7.8)
        else:
            entries = parse_entries(content)
            if entries:
                for entry in entries:
                    add_entry(doc, entry, master)
            else:
                for bullet in parse_publications(content):
                    add_bullet(doc, bullet, master)

    doc.save(out_path)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for filename, source in SOURCES.items():
        out_path = OUT_DIR / filename
        build_docx(source, out_path, master=(source.name == "yczresume.tex"))
        print(out_path)


if __name__ == "__main__":
    main()
